"""
Celery tasks for async report generation.
Reports are generated in the background to avoid blocking HTTP requests —
PDF/Excel generation can take several seconds for large datasets.
"""
import logging
import io
from celery import shared_task
from django.utils import timezone

logger = logging.getLogger(__name__)


@shared_task(bind=True, max_retries=3)
def generate_report_task(self, report_id: str):
    """
    Async report generation task.
    Generates PDF or Excel file, uploads to Supabase, updates Report.file_url.
    """
    from .models import Report
    try:
        report = Report.objects.select_related('tenant', 'site').get(id=report_id)
    except Report.DoesNotExist:
        logger.error('generate_report_task: report %s not found', report_id)
        return

    try:
        report.status = 'generating'
        report.save(update_fields=['status'])

        if report.format == 'pdf':
            file_url = _generate_pdf(report)
        else:
            file_url = _generate_excel(report)

        report.file_url = file_url
        report.status = 'ready'
        report.save(update_fields=['file_url', 'status'])
        logger.info('Report %s generated: %s', report_id, file_url)

    except Exception as exc:
        logger.exception('Report generation failed for %s: %s', report_id, exc)
        report.status = 'failed'
        report.save(update_fields=['status'])
        raise self.retry(exc=exc, countdown=60)


def _generate_pdf(report) -> str:
    """
    Build a PDF using ReportLab and upload to Supabase storage.
    Returns the public URL of the uploaded file.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    # Title
    story.append(Paragraph(
        f"ConstructTrack — {report.get_report_type_display()}",
        styles['Title']
    ))
    story.append(Paragraph(
        f"Tenant: {report.tenant.company_name} | Period: {report.date_from} to {report.date_to}",
        styles['Normal']
    ))
    story.append(Spacer(1, 12))

    # Content based on report type
    if report.report_type == 'material_consumption':
        story.extend(_material_report_content(report, styles))
    elif report.report_type == 'attendance_log':
        story.extend(_attendance_report_content(report, styles))
    elif report.report_type == 'monthly_spend':
        story.extend(_spend_report_content(report, styles))
    else:
        story.append(Paragraph('Report content will be generated here.', styles['Normal']))

    doc.build(story)
    buffer.seek(0)

    return _upload_to_supabase(buffer.read(), f"reports/{report.id}.pdf", 'application/pdf')


def _generate_excel(report) -> str:
    """Generate Excel report using openpyxl and upload to Supabase."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = report.get_report_type_display()

    ws['A1'] = f"ConstructTrack — {report.get_report_type_display()}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"Tenant: {report.tenant.company_name}"
    ws['A3'] = f"Period: {report.date_from} to {report.date_to}"

    if report.report_type == 'material_consumption':
        _excel_material_report(ws, report)
    elif report.report_type == 'attendance_log':
        _excel_attendance_report(ws, report)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return _upload_to_supabase(
        buffer.read(),
        f"reports/{report.id}.xlsx",
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


def _material_report_content(report, styles):
    """Build ReportLab content for material consumption report."""
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors
    from apps.bills.models import Bill

    site_filter = {'site': report.site} if report.site else {'site__tenant': report.tenant}
    bills = Bill.objects.filter(
        **site_filter,
        log_date__gte=report.date_from,
        log_date__lte=report.date_to,
    ).values('material_type', 'quantity', 'unit', 'total_amount_lkr', 'supplier_name', 'log_date')

    data = [['Date', 'Supplier', 'Material', 'Qty', 'Unit', 'Amount (LKR)']]
    for b in bills:
        data.append([
            str(b['log_date']), b['supplier_name'], b['material_type'],
            str(b['quantity']), b['unit'], str(b['total_amount_lkr']),
        ])

    if len(data) == 1:
        from reportlab.platypus import Paragraph
        return [Paragraph('No material records found for this period.', styles['Normal'])]

    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    return [table]


def _attendance_report_content(report, styles):
    from reportlab.platypus import Table, TableStyle, Paragraph
    from reportlab.lib import colors
    from apps.attendance.models import DailyAttendance

    site_filter = {'site': report.site} if report.site else {'site__tenant': report.tenant}
    records = DailyAttendance.objects.filter(
        **site_filter,
        log_date__gte=report.date_from,
        log_date__lte=report.date_to,
    ).select_related('worker')

    data = [['Date', 'Worker', 'Role', 'Status', 'Rate (LKR)', 'Earned (LKR)']]
    for r in records:
        data.append([
            str(r.log_date), r.worker.full_name, r.worker.role,
            r.status, str(r.daily_rate_lkr), str(r.total_earned_lkr),
        ])

    if len(data) == 1:
        return [Paragraph('No attendance records found for this period.', styles['Normal'])]

    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    return [table]


def _spend_report_content(report, styles):
    from reportlab.platypus import Paragraph
    from apps.bills.models import Bill
    from apps.attendance.models import DailyAttendance
    from django.db.models import Sum

    site_filter = {'site': report.site} if report.site else {'site__tenant': report.tenant}
    mat = Bill.objects.filter(**site_filter, log_date__gte=report.date_from, log_date__lte=report.date_to).aggregate(t=Sum('total_amount_lkr'))['t'] or 0
    wag = DailyAttendance.objects.filter(**site_filter, log_date__gte=report.date_from, log_date__lte=report.date_to).aggregate(t=Sum('total_earned_lkr'))['t'] or 0

    return [
        Paragraph(f"Total Material Spend: LKR {mat:,.2f}", styles['Normal']),
        Paragraph(f"Total Wage Spend: LKR {wag:,.2f}", styles['Normal']),
        Paragraph(f"Grand Total: LKR {float(mat) + float(wag):,.2f}", styles['Heading2']),
    ]


def _excel_material_report(ws, report):
    from apps.bills.models import Bill
    ws.append(['Date', 'Supplier', 'Material', 'Qty', 'Unit', 'Amount (LKR)'])
    site_filter = {'site': report.site} if report.site else {'site__tenant': report.tenant}
    bills = Bill.objects.filter(**site_filter, log_date__gte=report.date_from, log_date__lte=report.date_to)
    for b in bills:
        ws.append([str(b.log_date), b.supplier_name, b.material_type, float(b.quantity), b.unit, float(b.total_amount_lkr)])


def _excel_attendance_report(ws, report):
    from apps.attendance.models import DailyAttendance
    ws.append(['Date', 'Worker', 'Status', 'Rate', 'Earned'])
    site_filter = {'site': report.site} if report.site else {'site__tenant': report.tenant}
    records = DailyAttendance.objects.filter(**site_filter, log_date__gte=report.date_from, log_date__lte=report.date_to).select_related('worker')
    for r in records:
        ws.append([str(r.log_date), r.worker.full_name, r.status, float(r.daily_rate_lkr), float(r.total_earned_lkr)])


def _upload_to_supabase(file_bytes: bytes, path: str, content_type: str) -> str:
    """
    Upload file bytes to Supabase Storage and return the public URL.
    In development, if Supabase is not configured, saves locally and returns a placeholder URL.
    """
    from django.conf import settings
    import requests

    supabase_url = settings.SUPABASE_URL
    service_key = settings.SUPABASE_SERVICE_KEY
    bucket = settings.SUPABASE_BUCKET_NAME

    if not supabase_url or not service_key:
        # Development fallback — save to MEDIA_ROOT
        import os
        from pathlib import Path
        media_root = Path(settings.MEDIA_ROOT)
        file_path = media_root / path
        file_path.parent.mkdir(parents=True, exist_ok=True)
        with open(file_path, 'wb') as f:
            f.write(file_bytes)
        return f"{settings.MEDIA_URL}{path}"

    upload_url = f"{supabase_url}/storage/v1/object/{bucket}/{path}"
    response = requests.put(
        upload_url,
        data=file_bytes,
        headers={
            'Authorization': f'Bearer {service_key}',
            'Content-Type': content_type,
        },
    )
    response.raise_for_status()
    return f"{supabase_url}/storage/v1/object/public/{bucket}/{path}"
